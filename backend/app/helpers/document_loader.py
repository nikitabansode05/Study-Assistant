from pathlib import Path
import io
import fitz
from docx import Document
from PIL import Image
import pytesseract

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".markdown", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}

async def extract_file(path: Path) -> list[dict]:
    """Return text-bearing page/section records with source metadata."""
    ext = path.suffix.lower()
    output = []

    if ext == ".pdf":
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                # Attempt OCR for image/scanned pages that contain almost no extractable text.
                if len(text) < 30:
                    try:
                        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        image = Image.open(io.BytesIO(pix.tobytes("png")))
                        text = pytesseract.image_to_string(image).strip()
                    except Exception:
                        # Keep normal PDF extraction available even if Tesseract is absent.
                        pass
                if text:
                    output.append({"text": text, "page": i})

    elif ext == ".docx":
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text.strip() for cell in row.cells)
        if text.strip():
            output.append({"text": text.strip(), "page": None})

    elif ext in {".txt", ".md", ".markdown"}:
        # Markdown is indexed as plain text; headings, lists, tables, and code blocks are retained.
        text = path.read_text(encoding="utf-8-sig", errors="replace").strip()
        if text:
            output.append({"text": text, "page": None})

    elif ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        image = Image.open(path)
        text = pytesseract.image_to_string(image).strip()
        if text:
            output.append({"text": text, "page": None})
            
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not output:
        raise ValueError(
            "No text could be extracted. If this is a scan/image, install Tesseract OCR and ensure it is on PATH."
        )
    return output
